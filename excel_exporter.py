"""
Excel export functionality for enhanced results processing.

This module provides the ExcelExporter class for converting structured data
from analysis results into Excel files with proper formatting and templates.
"""

import os
from datetime import datetime
from typing import Dict, List, Optional, Any, Union
from dataclasses import dataclass
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from data_models import (
    ProcessedResult, StructuredData, DataTable, NamedEntity, 
    NumericValue, TemporalValue, EntityType
)


@dataclass
class ExportTemplate:
    """Template configuration for Excel exports."""
    name: str
    description: str
    include_summary: bool = True
    include_tables: bool = True
    include_entities: bool = True
    include_numeric_data: bool = True
    include_temporal_data: bool = True
    sheet_names: Dict[str, str] = None
    
    def __post_init__(self):
        if self.sheet_names is None:
            self.sheet_names = {
                'summary': 'Zusammenfassung',
                'tables': 'Tabellen',
                'entities': 'Entitäten',
                'numeric': 'Numerische Daten',
                'temporal': 'Zeitdaten'
            }


class ExcelExporter:
    """
    Handles export of structured data to Excel files with formatting and templates.
    """
    
    def __init__(self):
        self.templates = self._create_default_templates()
    
    def _create_default_templates(self) -> Dict[str, ExportTemplate]:
        """Create default export templates."""
        return {
            'complete': ExportTemplate(
                name='Vollständiger Export',
                description='Exportiert alle verfügbaren Daten in separate Arbeitsblätter'
            ),
            'tables_only': ExportTemplate(
                name='Nur Tabellen',
                description='Exportiert nur Tabellendaten',
                include_entities=False,
                include_numeric_data=False,
                include_temporal_data=False
            ),
            'summary': ExportTemplate(
                name='Zusammenfassung',
                description='Exportiert eine kompakte Zusammenfassung aller Daten',
                include_tables=False
            )
        }
    
    def export_to_excel(
        self, 
        result: ProcessedResult, 
        file_path: str,
        template_name: str = 'complete'
    ) -> bool:
        """
        Export a ProcessedResult to Excel file.
        
        Args:
            result: The ProcessedResult to export
            file_path: Path where to save the Excel file
            template_name: Name of the template to use
            
        Returns:
            bool: True if export was successful, False otherwise
        """
        try:
            template = self.templates.get(template_name, self.templates['complete'])
            workbook = Workbook()
            
            # Remove default sheet
            workbook.remove(workbook.active)
            
            # Create sheets based on template
            if template.include_summary:
                self._create_summary_sheet(workbook, result, template)
            
            if template.include_tables and result.extracted_data.tables:
                self._create_tables_sheet(workbook, result.extracted_data.tables, template)
            
            if template.include_entities and result.extracted_data.entities:
                self._create_entities_sheet(workbook, result.extracted_data.entities, template)
            
            if template.include_numeric_data and result.extracted_data.numeric_values:
                self._create_numeric_sheet(workbook, result.extracted_data.numeric_values, template)
            
            if template.include_temporal_data and result.extracted_data.temporal_data:
                self._create_temporal_sheet(workbook, result.extracted_data.temporal_data, template)
            
            # Ensure at least one sheet exists
            if not workbook.worksheets:
                self._create_summary_sheet(workbook, result, template)
            
            # Save the workbook
            workbook.save(file_path)
            return True
            
        except Exception as e:
            print(f"Fehler beim Excel-Export: {e}")
            return False
    
    def _create_summary_sheet(self, workbook: Workbook, result: ProcessedResult, template: ExportTemplate):
        """Create summary sheet with overview information."""
        sheet = workbook.create_sheet(template.sheet_names['summary'])
        
        # Header styling
        header_font = Font(bold=True, size=14)
        subheader_font = Font(bold=True, size=12)
        
        # Title
        sheet['A1'] = 'Analyse-Ergebnis Zusammenfassung'
        sheet['A1'].font = header_font
        
        row = 3
        
        # Basic information
        sheet[f'A{row}'] = 'Grundinformationen'
        sheet[f'A{row}'].font = subheader_font
        row += 1
        
        sheet[f'A{row}'] = 'Ergebnis-ID:'
        sheet[f'B{row}'] = result.id
        row += 1
        
        sheet[f'A{row}'] = 'Erstellt am:'
        sheet[f'B{row}'] = result.created_at.strftime('%d.%m.%Y %H:%M')
        row += 1
        
        sheet[f'A{row}'] = 'Analyse-Typ:'
        sheet[f'B{row}'] = result.metadata.analysis_type
        row += 1
        
        if result.source_info:
            sheet[f'A{row}'] = 'Quelle:'
            sheet[f'B{row}'] = result.source_info.type
            row += 1
            
            if result.source_info.file_name:
                sheet[f'A{row}'] = 'Dateiname:'
                sheet[f'B{row}'] = result.source_info.file_name
                row += 1
        
        row += 1
        
        # Data statistics
        sheet[f'A{row}'] = 'Datenstatistiken'
        sheet[f'A{row}'].font = subheader_font
        row += 1
        
        sheet[f'A{row}'] = 'Anzahl Tabellen:'
        sheet[f'B{row}'] = len(result.extracted_data.tables)
        row += 1
        
        sheet[f'A{row}'] = 'Anzahl Entitäten:'
        sheet[f'B{row}'] = len(result.extracted_data.entities)
        row += 1
        
        sheet[f'A{row}'] = 'Anzahl numerische Werte:'
        sheet[f'B{row}'] = len(result.extracted_data.numeric_values)
        row += 1
        
        sheet[f'A{row}'] = 'Anzahl Zeitdaten:'
        sheet[f'B{row}'] = len(result.extracted_data.temporal_data)
        row += 1
        
        # Content preview
        if result.content:
            row += 1
            sheet[f'A{row}'] = 'Inhalt (Vorschau)'
            sheet[f'A{row}'].font = subheader_font
            row += 1
            
            # Limit content to first 500 characters
            preview = result.content[:500]
            if len(result.content) > 500:
                preview += "..."
            
            sheet[f'A{row}'] = preview
            sheet.merge_cells(f'A{row}:D{row + 5}')
        
        # Auto-adjust column widths
        self._auto_adjust_columns(sheet)
    
    def _create_tables_sheet(self, workbook: Workbook, tables: List[DataTable], template: ExportTemplate):
        """Create sheet with all tables."""
        sheet = workbook.create_sheet(template.sheet_names['tables'])
        
        header_font = Font(bold=True, size=12)
        table_header_font = Font(bold=True, size=10, color="FFFFFF")
        table_header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        
        row = 1
        
        for i, table in enumerate(tables):
            # Table title
            title = table.title or f'Tabelle {i + 1}'
            sheet[f'A{row}'] = title
            sheet[f'A{row}'].font = header_font
            row += 2
            
            # Table headers
            for col, header in enumerate(table.headers, 1):
                cell = sheet.cell(row=row, column=col, value=header)
                cell.font = table_header_font
                cell.fill = table_header_fill
                cell.alignment = Alignment(horizontal='center')
            
            row += 1
            
            # Table data
            for table_row in table.rows:
                for col, value in enumerate(table_row, 1):
                    sheet.cell(row=row, column=col, value=value)
                row += 1
            
            row += 2  # Space between tables
        
        self._auto_adjust_columns(sheet)
    
    def _create_entities_sheet(self, workbook: Workbook, entities: List[NamedEntity], template: ExportTemplate):
        """Create sheet with extracted entities."""
        sheet = workbook.create_sheet(template.sheet_names['entities'])
        
        # Headers
        headers = ['Text', 'Typ', 'Konfidenz', 'Position Start', 'Position Ende']
        self._write_headers(sheet, headers, 1)
        
        # Data
        for row, entity in enumerate(entities, 2):
            sheet[f'A{row}'] = entity.text
            sheet[f'B{row}'] = entity.entity_type.value
            sheet[f'C{row}'] = f"{entity.confidence:.2%}"
            sheet[f'D{row}'] = entity.start_pos or ''
            sheet[f'E{row}'] = entity.end_pos or ''
        
        self._auto_adjust_columns(sheet)
    
    def _create_numeric_sheet(self, workbook: Workbook, numeric_values: List[NumericValue], template: ExportTemplate):
        """Create sheet with numeric values."""
        sheet = workbook.create_sheet(template.sheet_names['numeric'])
        
        # Headers
        headers = ['Wert', 'Einheit', 'Typ', 'Kontext']
        self._write_headers(sheet, headers, 1)
        
        # Data
        for row, num_val in enumerate(numeric_values, 2):
            sheet[f'A{row}'] = num_val.value
            sheet[f'B{row}'] = num_val.unit or ''
            sheet[f'C{row}'] = num_val.value_type or ''
            sheet[f'D{row}'] = num_val.context or ''
        
        self._auto_adjust_columns(sheet)
    
    def _create_temporal_sheet(self, workbook: Workbook, temporal_data: List[TemporalValue], template: ExportTemplate):
        """Create sheet with temporal data."""
        sheet = workbook.create_sheet(template.sheet_names['temporal'])
        
        # Headers
        headers = ['Datum/Zeit', 'Original Text', 'Präzision']
        self._write_headers(sheet, headers, 1)
        
        # Data
        for row, temp_val in enumerate(temporal_data, 2):
            sheet[f'A{row}'] = temp_val.value.strftime('%d.%m.%Y %H:%M:%S')
            sheet[f'B{row}'] = temp_val.original_text
            sheet[f'C{row}'] = temp_val.precision
        
        self._auto_adjust_columns(sheet)
    
    def _write_headers(self, sheet: Worksheet, headers: List[str], row: int):
        """Write formatted headers to a sheet."""
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=row, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
    
    def _auto_adjust_columns(self, sheet: Worksheet):
        """Auto-adjust column widths based on content."""
        for column in sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
            sheet.column_dimensions[column_letter].width = adjusted_width
    
    def export_structured_data_to_excel(
        self, 
        structured_data: StructuredData, 
        file_path: str,
        template_name: str = 'complete'
    ) -> bool:
        """
        Export StructuredData directly to Excel.
        
        Args:
            structured_data: The StructuredData to export
            file_path: Path where to save the Excel file
            template_name: Name of the template to use
            
        Returns:
            bool: True if export was successful, False otherwise
        """
        # Create a minimal ProcessedResult for compatibility
        from data_models import ProcessedResult, ResultMetadata
        
        result = ProcessedResult(
            content="Direkt exportierte strukturierte Daten",
            extracted_data=structured_data,
            metadata=ResultMetadata(analysis_type="Datenexport")
        )
        
        return self.export_to_excel(result, file_path, template_name)
    
    def get_available_templates(self) -> Dict[str, str]:
        """Get available export templates with descriptions."""
        return {name: template.description for name, template in self.templates.items()}
    
    def add_custom_template(self, name: str, template: ExportTemplate):
        """Add a custom export template."""
        self.templates[name] = template
    
    def preview_export_structure(self, result: ProcessedResult, template_name: str = 'complete') -> Dict[str, Any]:
        """
        Preview what would be exported without creating the file.
        
        Args:
            result: The ProcessedResult to preview
            template_name: Name of the template to use
            
        Returns:
            Dict with preview information
        """
        template = self.templates.get(template_name, self.templates['complete'])
        
        preview = {
            'template': template.name,
            'sheets': [],
            'total_rows': 0,
            'data_summary': {}
        }
        
        if template.include_summary:
            preview['sheets'].append(template.sheet_names['summary'])
            preview['total_rows'] += 20  # Estimated rows for summary
        
        if template.include_tables and result.extracted_data.tables:
            preview['sheets'].append(template.sheet_names['tables'])
            table_rows = sum(len(table.rows) + 3 for table in result.extracted_data.tables)
            preview['total_rows'] += table_rows
            preview['data_summary']['tables'] = len(result.extracted_data.tables)
        
        if template.include_entities and result.extracted_data.entities:
            preview['sheets'].append(template.sheet_names['entities'])
            preview['total_rows'] += len(result.extracted_data.entities) + 1
            preview['data_summary']['entities'] = len(result.extracted_data.entities)
        
        if template.include_numeric_data and result.extracted_data.numeric_values:
            preview['sheets'].append(template.sheet_names['numeric'])
            preview['total_rows'] += len(result.extracted_data.numeric_values) + 1
            preview['data_summary']['numeric_values'] = len(result.extracted_data.numeric_values)
        
        if template.include_temporal_data and result.extracted_data.temporal_data:
            preview['sheets'].append(template.sheet_names['temporal'])
            preview['total_rows'] += len(result.extracted_data.temporal_data) + 1
            preview['data_summary']['temporal_data'] = len(result.extracted_data.temporal_data)
        
        return preview


def create_excel_export_filename(result: ProcessedResult, base_name: str = None) -> str:
    """
    Create a standardized filename for Excel exports.
    
    Args:
        result: The ProcessedResult being exported
        base_name: Optional base name for the file
        
    Returns:
        str: Formatted filename with timestamp
    """
    if base_name is None:
        base_name = "analyse_ergebnis"
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    analysis_type = result.metadata.analysis_type.replace(" ", "_").lower()
    
    return f"{base_name}_{analysis_type}_{timestamp}.xlsx"


def validate_export_data(result: ProcessedResult) -> Dict[str, Any]:
    """
    Validate that a ProcessedResult has exportable data.
    
    Args:
        result: The ProcessedResult to validate
        
    Returns:
        Dict with validation results
    """
    validation = {
        'has_exportable_data': False,
        'warnings': [],
        'data_types': []
    }
    
    if result.extracted_data.tables:
        validation['has_exportable_data'] = True
        validation['data_types'].append('tables')
    
    if result.extracted_data.entities:
        validation['has_exportable_data'] = True
        validation['data_types'].append('entities')
    
    if result.extracted_data.numeric_values:
        validation['has_exportable_data'] = True
        validation['data_types'].append('numeric_values')
    
    if result.extracted_data.temporal_data:
        validation['has_exportable_data'] = True
        validation['data_types'].append('temporal_data')
    
    if not validation['has_exportable_data']:
        validation['warnings'].append('Keine strukturierten Daten zum Exportieren gefunden')
    
    if not result.content.strip():
        validation['warnings'].append('Kein Textinhalt vorhanden')
    
    return validation